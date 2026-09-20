"""app.services.report_exporter — Team Notes (CaseComment) and Investigator
Journal (CaseNote) sections in the PDF/HTML/DOCX exports. Neither was
previously pulled into any exported report format, despite both already
being distinct, named features elsewhere in the app."""
import io
import json
from app.services import report_exporter


def _make_case_with_investigation(app, user):
    from app.repositories.case_repository import create_case
    from app.repositories.investigation_repository import create_investigation
    with app.app_context():
        case = create_case("Report Notes Test Case", "desc", owner_user_id=user.id)
        inv = create_investigation(
            kind="domain", query="example.com",
            result_json=json.dumps({"domain": "example.com"}),
            user_id=user.id, case_id=case.id, confidence="CONFIRMED",
        )
    return case, [inv]


def _seed_comment_and_note(app, case, user):
    from app.repositories.case_comment_repository import add_comment
    from app.repositories.case_note_repository import add_note
    with app.app_context():
        add_comment(case.id, user.id, user.username, "This is a team comment about the case.")
        add_note(case.id, user.id, user.username, "lead", "Followed up on this lead, seems promising.")


# ── HTML ──────────────────────────────────────────────────────────────────

def test_html_includes_team_notes_and_journal_sections(app, user_a):
    case, invs = _make_case_with_investigation(app, user_a)
    _seed_comment_and_note(app, case, user_a)
    with app.app_context():
        html = report_exporter.export_html(case, invs, investigator=user_a.username)

    assert "Team Notes" in html
    assert "Investigator Journal" in html
    assert "This is a team comment about the case." in html
    assert "Followed up on this lead, seems promising." in html
    assert "LEAD" in html  # note kind, uppercased


def test_html_shows_placeholder_when_no_comments_or_notes(app, user_a):
    case, invs = _make_case_with_investigation(app, user_a)
    with app.app_context():
        html = report_exporter.export_html(case, invs, investigator=user_a.username)

    assert "No team notes recorded." in html
    assert "No investigator journal entries recorded." in html


# ── PDF (bytes only — just confirm it doesn't crash and grows with content) ──

def test_pdf_generates_with_and_without_notes(app, user_a):
    case, invs = _make_case_with_investigation(app, user_a)
    with app.app_context():
        empty_pdf = report_exporter.export_pdf(case, invs, investigator=user_a.username)
    _seed_comment_and_note(app, case, user_a)
    with app.app_context():
        filled_pdf = report_exporter.export_pdf(case, invs, investigator=user_a.username)

    assert isinstance(empty_pdf, (bytes, bytearray))
    assert isinstance(filled_pdf, (bytes, bytearray))
    assert len(filled_pdf) > len(empty_pdf)


# ── DOCX ──────────────────────────────────────────────────────────────────

def test_docx_includes_team_notes_and_journal_sections(app, user_a):
    from docx import Document

    case, invs = _make_case_with_investigation(app, user_a)
    _seed_comment_and_note(app, case, user_a)
    with app.app_context():
        docx_bytes = report_exporter.export_docx(case, invs, investigator=user_a.username)

    doc = Document(io.BytesIO(docx_bytes))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "TEAM NOTES" in full_text
    assert "INVESTIGATOR JOURNAL" in full_text
    assert "This is a team comment about the case." in full_text
    assert "Followed up on this lead, seems promising." in full_text


def test_docx_shows_none_recorded_placeholder(app, user_a):
    from docx import Document

    case, invs = _make_case_with_investigation(app, user_a)
    with app.app_context():
        docx_bytes = report_exporter.export_docx(case, invs, investigator=user_a.username)

    doc = Document(io.BytesIO(docx_bytes))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert full_text.count("None recorded.") == 2
